# Copyright (c) 2026, Muhammad Yusuf Tri Daryanto
# License: MIT

# W25/W26 — Import Stock Reconciliation via template opname yang ramah.
#
# Masalah: template Data Import native (parent/child row, kolom teknis, tanpa
# prefill stok) tidak layak pakai untuk tim gudang. Fitur ini: download template
# SATU BARIS PER ITEM yang sudah terisi stok saat ini dalam UOM inventaris item
# (W21/W23), tim mengisi kolom "Counted Qty" offline, lalu upload file dari
# FORM Stock Reconciliation (sejak W26; versi list-view W25 dipensiunkan —
# lihat retire_client_scripts di upgrade.py): endpoint TIDAK membuat dokumen
# apa pun, ia mengembalikan baris grid yang sudah dihitung penuh dan klien
# mengisi tabel items dokumennya sendiri (review + Submit tetap manual).
#
# Keputusan desain (plan W25 disetujui user + ruling advisor 2026-09-29):
#   - Custom flat importer, BUKAN reuse Data Import (formatnya tetap musuh
#     pengguna; hak aksesnya pun teritori System/Import Manager).
#   - Konversi UOM tetap DIJAGA hook W23 (apply_sr_inventory_uom,
#     before_validate) saat draft disimpan. Endpoint menghitung qty/rate native
#     untuk grid dengan MATEMATIKA YANG SAMA, karena fill grid programatik
#     (add_child + $.extend + refresh_field) TIDAK menjalankan handler baris
#     klien (custom_uom/custom_qty_after dst) — baris tanpa kolom turunan akan
#     menyimpan nilai mentah.
#   - Multicompany-safe: company dari DOKUMEN — dialog memfilter Warehouse ke
#     doc.company dan validasi menolak gudang milik perusahaan lain.
#   - Kolom "Current Qty" di template = snapshot info saat download; saldo
#     "before" saat upload selalu dihitung ulang dari ledger (get_previous_sle)
#     dengan posting date/time DOKUMEN, jadi SR backdated tetap benar.
#   - Draft-only: titik komit tetap tombol Submit manusia.
#   - Atomik: semua baris divalidasi dulu; error bernomor baris spreadsheet.
#   - Baris "counted == saldo ledger & rate kosong" di-strip dari hasil —
#     native remove_items_with_no_change akan membuangnya juga saat simpan;
#     tanpa strip, file yang seluruhnya cocok jadi EmptyStockReconciliation.
#
# Alur upload W26: klien pakai frappe.ui.FileUploader(as_dataurl) — v16
# mengembalikan data URL via on_success TANPA POST /api/method/upload_file
# (FileUploader tak bisa membawa konteks dokumen di FormData-nya) — lalu klien
# POST base64 itu sendiri ke endpoint ini.
#
# Izin: kedua endpoint mengecek has_permission("Stock Reconciliation",
# "create") — mengikuti semantik permission native (di site ini pemegang
# create adalah Stock Manager).

import base64
import binascii
import csv
import io
import re

import frappe
from frappe import _
from frappe.utils import cint, cstr, escape_html, flt, nowdate, nowtime

from warehouse_app.inventory_uom import (
	ITEM_UOM_FIELD,
	SR_AFTER_FIELD,
	SR_BEFORE_FIELD,
	SR_DIFF_FIELD,
	SR_FACTOR_FIELD,
	SR_RATE_FIELD,
	SR_UOM_FIELD,
)

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_DATA_ROWS = 5000
MAX_ERRORS_SHOWN = 30

H_ITEM = "Item Code"
H_ITEM_NAME = "Item Name"
H_WAREHOUSE = "Warehouse"
H_UOM = "UOM"
H_CURRENT = "Current Qty (as per UOM)"
H_COUNTED = "Counted Qty (as per UOM)"
H_RATE = "Valuation Rate (as per UOM)"

TEMPLATE_HEADERS = (H_ITEM, H_ITEM_NAME, H_WAREHOUSE, H_UOM, H_CURRENT, H_COUNTED, H_RATE)

# header (tanpa sufiks "(as per uom)", case-insensitive) -> key internal
_HEADER_KEYS = {
	"item code": "item_code",
	"item name": "item_name",
	"warehouse": "warehouse",
	"uom": "uom",
	"current qty": "current_qty",
	"counted qty": "counted",
	"valuation rate": "rate",
}


# ---------------------------------------------------------------------------
# Template (download)
# ---------------------------------------------------------------------------


def build_count_rows(warehouse, include_zero_stock=False):
	"""Baris prefill template: satu baris per item yang pernah bergerak di
	`warehouse` (default hanya yang stoknya ≠ 0). Item disabled dan item
	ber-serial/batch native dikecualikan (v1: hitung via form SR). UOM baris =
	Default Inventory UOM item (fallback stock UOM bila kosong/tak punya
	konversi); current qty = Bin.actual_qty ÷ faktor."""
	if not warehouse or not frappe.db.exists("Warehouse", warehouse):
		frappe.throw(_("Warehouse {0} not found.").format(frappe.bold(warehouse)))
	bin_filters = {"warehouse": warehouse}
	if not cint(include_zero_stock):
		bin_filters["actual_qty"] = ("!=", 0)
	bins = frappe.get_all(
		"Bin", filters=bin_filters, fields=["item_code", "actual_qty"], order_by="item_code"
	)
	if not bins:
		return []

	items = {
		d.name: d
		for d in frappe.get_all(
			"Item",
			filters={"name": ("in", [b.item_code for b in bins])},
			fields=[
				"name",
				"item_name",
				"stock_uom",
				"disabled",
				"has_batch_no",
				"has_serial_no",
				ITEM_UOM_FIELD,
			],
		)
	}
	default_uoms = {
		it.name: it.get(ITEM_UOM_FIELD)
		for it in items.values()
		if it.get(ITEM_UOM_FIELD) and it[ITEM_UOM_FIELD] != it.stock_uom
	}
	factors = {}
	if default_uoms:
		for d in frappe.get_all(
			"UOM Conversion Detail",
			filters={"parent": ("in", sorted(default_uoms))},
			fields=["parent", "uom", "conversion_factor"],
		):
			factors[(d.parent, d.uom)] = flt(d.conversion_factor)

	rows = []
	for b in bins:
		item = items.get(b.item_code)
		if not item or item.disabled or item.has_batch_no or item.has_serial_no:
			continue
		uom = item.get(ITEM_UOM_FIELD) or item.stock_uom
		factor = 1.0 if uom == item.stock_uom else factors.get((item.name, uom))
		if not factor:
			# Default Inventory UOM tak punya baris konversi → fallback stock UOM
			uom, factor = item.stock_uom, 1.0
		rows.append(
			{
				"item_code": item.name,
				"item_name": item.item_name or item.name,
				"warehouse": warehouse,
				"uom": uom,
				"current_qty": flt(flt(b.actual_qty) / factor, 3),
			}
		)
	rows.sort(key=lambda r: r["item_code"])
	if len(rows) > MAX_DATA_ROWS:
		# Simetri dengan cap parser upload — template yang lebih besar dari cap
		# hanya jadi jebakan (di-download, diisi, lalu ditolak saat upload).
		frappe.throw(
			_(
				"The warehouse has more than {0} countable items. The import supports at most {0} rows per file."
			).format(frappe.bold(MAX_DATA_ROWS))
		)
	return rows


def make_template_bytes(rows, file_type="xlsx"):
	"""rows (dict hasil build_count_rows) → byte file template. Satu sheet,
	row 1 = header (tanpa baris judul — parser memetakan berdasarkan nama
	header). Kolom Counted/Rate sengaja kosong (input user)."""
	data = [list(TEMPLATE_HEADERS)]
	for r in rows or []:
		data.append(
			[
				r.get("item_code"),
				r.get("item_name"),
				r.get("warehouse"),
				r.get("uom"),
				r.get("current_qty"),
				"",  # Counted Qty (as per UOM)
				"",  # Valuation Rate (as per UOM)
			]
		)
	if file_type == "csv":
		from frappe.utils.csvutils import to_csv

		# BOM utf-8-sig agar Excel Windows membaca UTF-8 dengan benar
		return ("\ufeff" + to_csv(data)).encode("utf-8")
	from frappe.utils.xlsxutils import make_xlsx

	return make_xlsx(data, _("Stock Count"), column_widths=[18, 40, 28, 10, 16, 16, 18]).getvalue()


@frappe.whitelist(methods=["POST"])
def download_stock_count_template(warehouse=None, include_zero_stock=0, file_type="xlsx"):
	frappe.has_permission("Stock Reconciliation", "create", throw=True)
	warehouse = cstr(warehouse).strip()
	if not warehouse:
		frappe.throw(_("Warehouse is required."))
	rows = build_count_rows(warehouse, cint(include_zero_stock))
	file_type = file_type if file_type in ("xlsx", "csv") else "xlsx"
	name = "Stock_Count_{0}_{1}".format(frappe.scrub(warehouse), nowdate().replace("-", ""))
	from frappe.desk.utils import provide_binary_file

	provide_binary_file(name, file_type, make_template_bytes(rows, file_type))


# ---------------------------------------------------------------------------
# Parser (upload)
# ---------------------------------------------------------------------------


def _norm_header(raw):
	return re.sub(r"\s+", " ", cstr(raw).strip()).lower().replace("(as per uom)", "").strip()


def _header_key(raw):
	return _HEADER_KEYS.get(_norm_header(raw))


def _headers_recognized(header):
	keys = {_header_key(cell) for cell in header}
	keys.discard(None)
	return "item_code" in keys and "counted" in keys


def _xlsx_rows(content):
	"""Baca sheet aktif openpyxl read_only (streaming) dengan cap baris SAAT
	iterasi — file salah-unggah raksasa gagal cepat, bukan setelah termuat."""
	from openpyxl import load_workbook

	try:
		wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
	except Exception:
		# zip rusak / file di-rename dari format lain → pesan ramah, bukan 500
		frappe.throw(
			_("The file is not a valid .xlsx. Re-download the template and edit that copy.")
		)
	try:
		rows = []
		for row in wb.active.iter_rows(values_only=True):
			rows.append(list(row))
			if len(rows) > MAX_DATA_ROWS + 1:
				frappe.throw(
					_("The file has more than {0} data rows.").format(frappe.bold(MAX_DATA_ROWS))
				)
		return rows
	finally:
		wb.close()


def _csv_rows(content):
	"""→ (rows, delimiter). Excel locale Indonesia menyimpan CSV ber-delimiter
	';'. Coba ',' dulu; bila headernya tak dikenali dan versi ';' dikenali,
	gunakan ';' — dialek menentukan cara membaca angka (lihat _parse_number)."""
	text = None
	for encoding in ("utf-8-sig", "utf-8", "cp1252"):
		try:
			text = content.decode(encoding)
			break
		except UnicodeDecodeError:
			continue
	if text is None:
		frappe.throw(_("The CSV file is not readable as UTF-8 or Windows-1252 text."))
	rows = list(csv.reader(io.StringIO(text), delimiter=","))
	if not rows or not _headers_recognized(rows[0]):
		semi_rows = list(csv.reader(io.StringIO(text), delimiter=";"))
		if semi_rows and _headers_recognized(semi_rows[0]):
			return semi_rows, ";"
	return rows, ","


_COMMA_DECIMAL = re.compile(r"^-?\d+,\d+$")
_DOT_GROUPING = re.compile(r"^-?\d{1,3}(?:\.\d{3})+$")
_PLAIN_NUMBER = re.compile(r"^-?\d+(?:\.\d+)?$")


def _parse_number(value, row_no, label, allow_grouping=False):
	"""→ (float | None, pesan_error | None). None tanpa error = sel kosong.
	Menerima angka type xlsx, teks desimal koma ("1,5"); grouping titik
	("1.500" = 1500) HANYA untuk dialek CSV ';' (Excel locale ID) — pada
	dialeks koma/xlsx titik selalu dibaca desimal ("1.500" = 1.5). Menolak
	campuran pemisah dan teks tak dikenali."""
	if value is None or (isinstance(value, str) and not value.strip()):
		return None, None
	if isinstance(value, bool) or not isinstance(value, (int, float, str)):
		return None, _("Row {0}: {1} must be a number.").format(row_no, label)
	if isinstance(value, (int, float)):
		return flt(value), None
	text = value.strip()
	if "," in text and "." in text:
		return None, _(
			'Row {0}: {1} value "{2}" mixes "," and ".". Use plain numbers (e.g. 1.5).'
		).format(row_no, label, escape_html(text))
	if _COMMA_DECIMAL.match(text):
		return flt(text.replace(",", ".")), None
	if allow_grouping and _DOT_GROUPING.match(text):
		return flt(text.replace(".", "")), None
	if _PLAIN_NUMBER.match(text):
		return flt(text), None
	return None, _("Row {0}: {1} value {2} is not a recognized number.").format(
		row_no, label, frappe.bold(escape_html(text))
	)


def parse_count_content(content, filename):
	"""Byte file + nama file → {"rows": [...], "skipped": n, "warnings": [...]}.
	Row ber-nomor baris spreadsheet (header = 1). Semua error struktural
	dikumpulkan lalu dilempar sekaligus (atomik)."""
	if len(content or b"") > MAX_FILE_BYTES:
		frappe.throw(
			_("The file is larger than {0} MB.").format(frappe.bold(MAX_FILE_BYTES // (1024 * 1024)))
		)
	ext = cstr(filename).lower().rsplit(".", 1)[-1] if "." in cstr(filename) else ""
	if ext == "xlsx":
		rows = _xlsx_rows(content)
		allow_grouping = False
	elif ext == "csv":
		rows, delimiter = _csv_rows(content)
		allow_grouping = delimiter == ";"
	else:
		frappe.throw(
			_("Only .xlsx and .csv count files are supported (got {0}).").format(
				frappe.bold(escape_html(ext or cstr(filename) or "?"))
			)
		)
	if not rows:
		frappe.throw(_("The file is empty. Please use the downloaded template."))
	if len(rows) > MAX_DATA_ROWS + 1:
		frappe.throw(_("The file has more than {0} data rows.").format(frappe.bold(MAX_DATA_ROWS)))

	colmap = {}
	unknown = []
	for idx, cell in enumerate(rows[0]):
		key = _header_key(cell)
		if key and key not in colmap.values():
			colmap[idx] = key  # kemunculan pertama menang
		elif _norm_header(cell) and not key:
			unknown.append(_norm_header(cell))

	def cell_of(raw, key):
		for idx, k in colmap.items():
			if k == key:
				return raw[idx] if idx < len(raw) else None
		return None

	missing = [
		label
		for label, key in (
			(H_ITEM, "item_code"),
			(H_WAREHOUSE, "warehouse"),
			(H_UOM, "uom"),
			(H_COUNTED, "counted"),
		)
		if key not in colmap.values()
	]
	if missing:
		frappe.throw(
			_("Missing required column(s): {0}. Use the template downloaded from this dialog.").format(
				", ".join(frappe.bold(m) for m in missing)
			)
		)

	warnings = []
	if unknown:
		warnings.append(
			_("Ignored unknown column(s): {0}").format(
				", ".join(escape_html(u) for u in unknown)
			)
		)

	errors = []
	parsed_rows = []
	skipped = 0
	for row_no, raw in enumerate(rows[1:], start=2):
		if all(cstr(v).strip() == "" for v in raw):
			continue  # baris kosong sisa area sheet
		row_errors = []
		item_code = cstr(cell_of(raw, "item_code")).strip()
		warehouse = cstr(cell_of(raw, "warehouse")).strip()
		uom = cstr(cell_of(raw, "uom")).strip()
		counted, err = _parse_number(cell_of(raw, "counted"), row_no, H_COUNTED, allow_grouping)
		if err:
			row_errors.append(err)
		rate, err = _parse_number(cell_of(raw, "rate"), row_no, H_RATE, allow_grouping)
		if err:
			row_errors.append(err)
		if counted is None and not row_errors:
			if not item_code:
				continue  # serpihan tanpa item & tanpa angka — diloncati
			skipped += 1
			continue
		if not item_code:
			row_errors.append(
				_("Row {0}: Item Code is required when Counted Qty is filled.").format(row_no)
			)
		if not warehouse:
			row_errors.append(_("Row {0}: Warehouse is required.").format(row_no))
		if not uom:
			row_errors.append(
				_("Row {0}: UOM is required when Counted Qty is filled.").format(row_no)
			)
		if counted is not None and counted < 0:
			row_errors.append(
				_("Row {0}: Counted Qty cannot be negative (got {1}).").format(row_no, counted)
			)
		if rate is not None:
			if rate < 0:
				row_errors.append(
					_("Row {0}: Valuation Rate cannot be negative (got {1}).").format(row_no, rate)
				)
			if rate == 0:
				row_errors.append(
					_(
						"Row {0}: Valuation Rate 0 cannot be imported. Leave the cell blank to keep the system rate."
					).format(row_no)
				)
		if row_errors:
			errors.extend(row_errors)
			continue
		parsed_rows.append(
			{
				"row_no": row_no,
				"item_code": item_code,
				"warehouse": warehouse,
				"uom": uom,
				"counted": counted,
				"rate": rate,
			}
		)

	if errors:
		frappe.throw(
			"<br>".join(errors[:MAX_ERRORS_SHOWN])
			+ ("<br>…" if len(errors) > MAX_ERRORS_SHOWN else "")
		)
	if not parsed_rows:
		frappe.throw(_("No rows with a Counted Qty were found in the file."))
	return {"rows": parsed_rows, "skipped": skipped, "warnings": warnings}


# ---------------------------------------------------------------------------
# Validasi + baris grid (W26 — murni, endpoint TIDAK membuat dokumen)
# ---------------------------------------------------------------------------


def validate_count_rows(parsed, company, posting_date=None, posting_time=None):
	"""parsed (hasil parse_count_content, atau dict bentuk sama) → baris grid
	SIAP PAKAI untuk tabel items form Stock Reconciliation (W26). TIDAK membuat
	dokumen — klien mengisi grid-nya sendiri (review + Submit tetap manual).

	Validasi semantik semua baris dulu (error bernomor baris spreadsheet,
	atomik). `company` WAJIB dan menjadi acak keabsahan gudang
	(multicompany-safe): setiap gudang wajib milik company dokumen.

	Setiap kolom turunan (custom W23 + qty/valuation_rate native) dihitung
	server-side di sini — matematika mengikuti hook apply_sr_inventory_uom
	(satu jalur konversi yang teruji); saat draft disimpan, hook menghitung
	ulang nilai yang sama. Ledger dibaca pada posting_date/posting_time yang
	DIBERIKAN (mirror native get_items) supaya SR backdated tetap benar."""
	company = cstr(company).strip()
	if not company:
		frappe.throw(_("Company is required to import a stock count."))
	posting_date = posting_date or nowdate()
	posting_time = posting_time or nowtime()

	rows = [dict(r) for r in (parsed or {}).get("rows") or []]
	if not rows:
		frappe.throw(_("No counted rows to import."))
	warnings = list((parsed or {}).get("warnings") or [])
	errors = []

	items = {
		d.name: d
		for d in frappe.get_all(
			"Item",
			filters={"name": ("in", sorted({r["item_code"] for r in rows}))},
			fields=[
				"name",
				"item_name",
				"stock_uom",
				"disabled",
				"has_batch_no",
				"has_serial_no",
				ITEM_UOM_FIELD,
			],
		)
	}
	warehouses = {
		d.name: d
		for d in frappe.get_all(
			"Warehouse",
			filters={"name": ("in", sorted({r["warehouse"] for r in rows}))},
			fields=["name", "company", "is_group", "disabled"],
		)
	}
	conv_needed = {
		(r["item_code"], r["uom"])
		for r in rows
		if items.get(r["item_code"]) and r["uom"] != items[r["item_code"]].stock_uom
	}
	factors = {}
	if conv_needed:
		for d in frappe.get_all(
			"UOM Conversion Detail",
			filters={"parent": ("in", sorted({c[0] for c in conv_needed}))},
			fields=["parent", "uom", "conversion_factor"],
		):
			factors[(d.parent, d.uom)] = flt(d.conversion_factor)

	seen = set()
	# Semua nilai r[...] berasal dari file upload → escape sebelum masuk pesan
	# HTML (frappe.msgprint merender HTML di sisi klien).
	for r in rows:
		n = r["row_no"]
		item = items.get(r["item_code"])
		if not item:
			errors.append(
				_("Row {0}: Item {1} not found.").format(n, frappe.bold(escape_html(r["item_code"])))
			)
			continue
		if item.disabled:
			errors.append(_("Row {0}: Item {1} is disabled.").format(n, frappe.bold(item.name)))
			continue
		if item.has_batch_no or item.has_serial_no:
			errors.append(
				_(
					"Row {0}: Item {1} is serial/batch tracked. Count it via the Stock Reconciliation form."
				).format(n, frappe.bold(item.name))
			)
			continue
		wh = warehouses.get(r["warehouse"])
		if not wh:
			errors.append(
				_("Row {0}: Warehouse {1} not found.").format(
					n, frappe.bold(escape_html(r["warehouse"]))
				)
			)
			continue
		if wh.is_group or wh.disabled:
			errors.append(
				_("Row {0}: Warehouse {1} is a group or disabled warehouse.").format(
					n, frappe.bold(escape_html(r["warehouse"]))
				)
			)
			continue
		if not wh.company:
			errors.append(
				_("Row {0}: Warehouse {1} has no company set.").format(
					n, frappe.bold(escape_html(r["warehouse"]))
				)
			)
			continue
		if wh.company != company:
			errors.append(
				_(
					"Row {0}: Warehouse {1} belongs to company {2}, but this reconciliation is for {3}."
				).format(
					n,
					frappe.bold(escape_html(r["warehouse"])),
					frappe.bold(wh.company),
					frappe.bold(escape_html(company)),
				)
			)
			continue
		if not r["uom"]:
			errors.append(_("Row {0}: UOM is required when Counted Qty is filled.").format(n))
			continue
		factor = 1.0
		if r["uom"] != item.stock_uom:
			factor = factors.get((item.name, r["uom"]))
			if not factor:
				errors.append(
					_(
						"Row {0}: UOM {1} is not valid for Item {2}. Use the Stock UOM or add it to the item's UOM conversion table."
					).format(n, frappe.bold(escape_html(r["uom"])), frappe.bold(item.name))
				)
				continue
		key = (item.name, r["warehouse"])
		if key in seen:
			errors.append(
				_(
					"Row {0}: Item {1} with Warehouse {2} appears more than once in the file."
				).format(
					n, frappe.bold(item.name), frappe.bold(escape_html(r["warehouse"]))
				)
			)
			continue
		seen.add(key)
		if flt(r["counted"]) < 0:
			errors.append(
				_("Row {0}: Counted Qty cannot be negative (got {1}).").format(n, r["counted"])
			)
			continue
		if r["rate"] is not None:
			if flt(r["rate"]) < 0:
				errors.append(
					_("Row {0}: Valuation Rate cannot be negative (got {1}).").format(n, r["rate"])
				)
				continue
			if flt(r["rate"]) == 0:
				errors.append(
					_(
						"Row {0}: Valuation Rate 0 cannot be imported. Leave the cell blank to keep the system rate."
					).format(n)
				)
				continue
		r["_factor"] = factor

	if errors:
		frappe.throw(
			"<br>".join(errors[:MAX_ERRORS_SHOWN])
			+ ("<br>…" if len(errors) > MAX_ERRORS_SHOWN else "")
		)

	ledger_cache = {}

	def ledger_snapshot(item_code, warehouse):
		"""→ (qty_after_transaction, valuation_rate) pada posting timestamp —
		satu panggilan per item+gudang, kedua kolom dipakai sekaligus."""
		key = (item_code, warehouse)
		if key not in ledger_cache:
			from erpnext.stock.stock_ledger import get_previous_sle

			prev = get_previous_sle(
				{
					"item_code": item_code,
					"warehouse": warehouse,
					"posting_date": posting_date,
					"posting_time": posting_time,
				}
			) or {}
			ledger_cache[key] = (
				flt(prev.get("qty_after_transaction")),
				flt(prev.get("valuation_rate")),
			)
		return ledger_cache[key]

	kept, matched = [], 0
	for r in rows:
		current, system_rate = ledger_snapshot(r["item_code"], r["warehouse"])
		before = flt(current) / r["_factor"]
		# Toleransi 0.0005 (di bawah presisi 3 desimal kolom before): pembulatan
		# bisa membuat qty "berubah" versi parse ternyata identik versi native —
		# native remove_items_with_no_change akan membuangnya juga saat simpan.
		# Strip di sini supaya baris tsb tak pernah sampai ke grid; tanpa ini,
		# file yang seluruhnya cocok melempar EmptyStockReconciliationItemsError
		# mentah di form.
		if r["rate"] is None and abs(flt(r["counted"]) - flt(before)) < 0.0005:
			matched += 1
			continue
		if not current and flt(r["counted"]) > 0 and r["rate"] is None:
			warnings.append(
				_(
					"Row {0}: Item {1} has zero stock. Fill Valuation Rate (as per UOM) or it cannot be submitted."
				).format(r["row_no"], frappe.bold(escape_html(r["item_code"])))
			)
		r["_snapshot"] = (current, system_rate, before)
		kept.append(r)

	stats = {
		"company": company,
		"added": 0,
		"skipped": cint((parsed or {}).get("skipped")),
		"matched": matched,
		"warnings": warnings,
		"rows": [],
	}
	if not kept:
		stats["message"] = _("All counts match the current stock. Nothing to import.")
		return stats

	use_serial_batch_fields = cint(
		frappe.db.get_single_value("Stock Settings", "use_serial_batch_fields")
	)
	grid_rows = []
	for r in kept:
		current, system_rate, before = r["_snapshot"]
		factor = r["_factor"]
		counted = flt(r["counted"])
		row = {
			"item_code": r["item_code"],
			"item_name": items[r["item_code"]].item_name or r["item_code"],
			"warehouse": r["warehouse"],
			"stock_uom": items[r["item_code"]].stock_uom,
			SR_UOM_FIELD: r["uom"],
			SR_FACTOR_FIELD: factor,
			"current_qty": current,
			"current_valuation_rate": system_rate,
			SR_BEFORE_FIELD: flt(before, 3),
			SR_AFTER_FIELD: counted,
			SR_DIFF_FIELD: flt(counted - flt(before, 3), 3),
			# native qty (stock UOM) & valuation_rate — matematika hook W23.
			# Rate kosong = saldo ledger (prefill gaya native fetch); saldo-0 →
			# None (bukan 0 eksplisit) supaya native backfill saat submit,
			# bukan "Valuation Rate required".
			"qty": flt(counted * factor),
			"valuation_rate": flt(r["rate"] / factor) if r["rate"] is not None else (system_rate or None),
			"use_serial_batch_fields": use_serial_batch_fields,
		}
		if r["rate"] is not None:
			# Kolom review W23 ikut diisi bila file memberi rate eksplisit —
			# hook menghitung ulang rate/factor yang identik saat simpan
			# (rate ≠ 0 lolos guard zero-storage-nya).
			row[SR_RATE_FIELD] = r["rate"]
		grid_rows.append(row)

	stats["added"] = len(grid_rows)
	stats["rows"] = grid_rows
	return stats


@frappe.whitelist(methods=["POST"])
def upload_stock_count(filename=None, data=None, company=None, posting_date=None, posting_time=None):
	"""File count (base64 murni atau data URL dari FileUploader as_dataurl) +
	konteks dokumen (company, posting date/time) → dict baris grid untuk form
	SR. TIDAK membuat dokumen (W26)."""
	frappe.has_permission("Stock Reconciliation", "create", throw=True)
	filename = cstr(filename).strip()
	data = cstr(data).strip()
	if not filename or not data:
		frappe.throw(_("Please choose a .xlsx or .csv count file to upload."))
	# Data URL "data:<mime>;base64,<payload>" — buang prefix-nya; base64 murni
	# (curl/gate) tetap diterima apa adanya.
	if data[:5].lower() == "data:":
		data = data.split(",", 1)[1] if "," in data else ""
	# Fail-fast sebelum decode: base64 membengkak ~4/3 — tolak payload raksasa
	# agar tidak di-decode percuma (cap byte akhir tetap di parse_count_content).
	if len(data) > (MAX_FILE_BYTES * 4 // 3) + 1024:
		frappe.throw(
			_("The file is larger than {0} MB.").format(frappe.bold(MAX_FILE_BYTES // (1024 * 1024)))
		)
	try:
		content = base64.b64decode(data, validate=False)
	except binascii.Error:
		frappe.throw(
			_(
				"The file could not be decoded. Re-download the template and edit that copy."
			)
		)
	parsed = parse_count_content(content, filename)
	return validate_count_rows(parsed, company, posting_date, posting_time)
