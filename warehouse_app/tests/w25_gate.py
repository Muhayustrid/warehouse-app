# Copyright (c) 2026, Muhammad Yusuf Tri Daryanto
# License: MIT

# Gate W25/W26 — import opname Stock Reconciliation (template ter-prefill →
# upload CSV/Excel → baris grid terhitung penuh → klien mengisi form SR;
# konversi UOM dijaga hook W23).
#
# Jalankan:
#   docker exec erpnext-new-backend-1 bench --site frontend execute \
#       warehouse_app.tests.w25_gate.run_gate
#
# Kontrak output: TEPAT satu baris GATE_JSON:{...} (ok: bool, checks: dict,
# error: opsional). Cek gagal -> ok=false TANPA raise; hanya crash tak terduga
# yang di-raise (bench exit nonzero).
#
# Yang diverifikasi (nomor = urut eksekusi):
#   1. Client Script SR: list-view W25 PENSIUN (tak ada doc memuat markernya),
#      form-view W26 import (enabled, view Form, marker + role gate), dan
#      form-view W23 UOM tetap hidup berdampingan.
#   2. Modul sr_import + 2 endpoint whitelisted POST-only ber-cek
#      has_permission create SR + smoke decode base64/data-URL via
#      upload_stock_count (endpoint murni — tanpa insert).
#   3. build_count_rows: prefill UOM inventaris (Box) + current qty ÷ faktor
#      (24/12 = 2); stok-0 hanya muncul dengan toggle; item batch dikecualikan.
#   4. Roundtrip file: xlsx (edit sel Counted/Rate via openpyxl) + csv
#      (+ varian delimiter ";" & desimal koma "3,5" & grouping "1.500"); baris
#      tanpa Counted Qty terskip oleh parser (stat skipped, bukan row).
#   5. validate_count_rows: baris grid terhitung penuh (qty = counted × faktor
#      = 36, before 2, after 3, diff 1, rate 1200/12 = 100, current 24, rate
#      per UOM 1200) — lalu SR dibangun persis seperti klien dari baris tsb →
#      hook W23 mengonversi ke nilai sama (docstatus 0).
#   6. Semua baris cocok ledger → added=0, rows=[], pesan match, TANPA membuat
#      dokumen SR apa pun.
#   7. Error parse (ekstensi salah / kolom wajib hilang / rate 0 / angka
#      sampah / pemisah campuran) + error validate atomik bernomor baris
#      (item tak dikenal / duplikat / UOM kosong / negatif / rate 0 /
#      mismatch company) — murni, tak ada dokumen tersisa.
#   8. Permission: user tanpa role ditolak upload_stock_count.
#   9. Item stok-0 tanpa rate → warning "Valuation Rate", baris tetap
#      dikembalikan (draft disimpan klien).
#  10. Submit draft hasil import → SLE qty_after_transaction 36 @ 100.
#  11. upgrade.apply() idempoten dua kali: list W25 tetap absen, form W26
#      tetap ada + enabled.
#
# Fixture (prefix "ZZTEST-W25"): item1 Nos+Box×12 (default Box) disemai SE
# Material Receipt 24 Nos @100; item2 (stok 0, Bin programmatik) utk kasus
# warning; item3 (has_batch_no) utk eksklusi. Teardown di finally, residu
# prefix = 0 (termasuk user + Sessions/Activity Log).

import base64
import csv
import inspect
import io
import json
import traceback

import frappe
from frappe.utils import flt, nowdate, nowtime

from warehouse_app import upgrade
from warehouse_app.inventory_uom import (
	SR_AFTER_FIELD,
	SR_BEFORE_FIELD,
	SR_DIFF_FIELD,
	SR_FACTOR_FIELD,
	SR_RATE_FIELD,
	SR_UOM_FIELD,
)
from warehouse_app.tests.guard import count_residue
from warehouse_app.warehouse_app import sr_import

PREFIX = "ZZTEST-W25"
MARKER = upgrade.SR_LIST_SCRIPT_MARKER  # marker W25 yang DIPENSIUNKAN
TRACKED = {"sr": [], "se": [], "item": [], "user": []}


class GateAborted(Exception):
	"""Prasyarat gate gagal; sisa langkah dilewati, GATE_JSON tetap dicetak (ok=false)."""


def run_gate():
	checks = {}
	error = None

	def check(name, ok, evidence):
		checks[name] = str(evidence) if ok else f"GAGAL: {evidence}"

	try:
		_run_gate(check)
	except GateAborted:
		pass
	except Exception:
		error = traceback.format_exc()
	finally:
		try:
			_teardown(check)
		except Exception as te:
			check("teardown", False, f"teardown error {type(te).__name__}: {te}")
		payload = {
			"ok": error is None and not any(v.startswith("GAGAL") for v in checks.values()),
			"checks": checks,
		}
		if error:
			payload["error"] = error[-2000:]
		print("GATE_JSON:" + json.dumps(payload, ensure_ascii=False))

	if error:
		raise SystemExit(1)


def _make_item(stock, alt, item_group, name_suffix):
	item = frappe.get_doc(
		{
			"doctype": "Item",
			"item_code": PREFIX + "-" + frappe.generate_hash(length=8),
			"item_name": PREFIX + " " + name_suffix,
			"item_group": item_group,
			"stock_uom": stock,
			"is_stock_item": 1,
			"has_batch_no": 0,
			"has_serial_no": 0,
			"uoms": [{"uom": alt, "conversion_factor": 12}],
			upgrade.ITEM_UOM_FIELD: alt,
		}
	)
	item.insert()
	TRACKED["item"].append(item.name)
	return item


def _make_bin(item_code, warehouse):
	if frappe.db.exists("Bin", {"item_code": item_code, "warehouse": warehouse}):
		return
	doc = frappe.get_doc({"doctype": "Bin", "item_code": item_code, "warehouse": warehouse})
	doc.flags.ignore_permissions = True
	doc.insert()


def _seed_stock(check, item_code, warehouse, company):
	try:
		se = frappe.get_doc(
			{
				"doctype": "Stock Entry",
				"stock_entry_type": "Material Receipt",
				"company": company,
				"posting_date": nowdate(),
				"items": [{"item_code": item_code, "t_warehouse": warehouse, "qty": 24, "basic_rate": 100}],
			}
		)
		se.insert()
		se.submit()
		TRACKED["se"].append(se.name)
		frappe.db.commit()
		bin_qty = flt(
			frappe.db.get_value("Bin", {"item_code": item_code, "warehouse": warehouse}, "actual_qty")
		)
		ok = abs(bin_qty - 24) < 1e-6
		check("fixture_seed", ok, f"SE={se.name}, bin={bin_qty} (want 24)")
		return ok
	except Exception as e:
		check("fixture_seed", False, f"{type(e).__name__}: {str(e)[:200]}")
		return False


def _run_gate(check):
	check("pre_clean", *_sweep())

	# --- Preflight (environment-aware) ---
	stock, alt, invalid = _pick_uoms()
	item_group = _pick_item_group()
	company = "JURI" if frappe.db.exists("Company", "JURI") else None
	if not company:
		companies = frappe.get_all("Company", pluck="name", limit=1)
		company = companies[0] if companies else None
	warehouse = None
	if company:
		warehouses = frappe.get_all(
			"Warehouse",
			filters={"company": company, "is_group": 0, "disabled": 0},
			pluck="name",
			limit=1,
		)
		warehouse = warehouses[0] if warehouses else None
	check(
		"preflight",
		bool(stock and alt and item_group and company and warehouse),
		f"stock={stock!r}, alt={alt!r}, invalid={invalid!r}, group={item_group!r}, "
		f"company={company!r}, warehouse={warehouse!r}",
	)
	if not (stock and alt and item_group and company and warehouse):
		raise GateAborted()

	# --- 1. Client Script SR (W25 pensiun + W26 form import + W23 form UOM) ---
	try:
		list_left = frappe.get_all(
			"Client Script",
			filters={"script": ("like", "%" + MARKER + "%")},
			pluck="name",
		)
		w26 = frappe.db.get_value(
			"Client Script",
			"warehouse-app-w26-stock-count-import",
			["dt", "enabled", "view", "script"],
			as_dict=1,
		)
		form_script = frappe.db.get_value(
			"Client Script",
			{"dt": "Stock Reconciliation", "script": ("like", "%" + upgrade.SR_SCRIPT_MARKER + "%")},
			["enabled", "view"],
			as_dict=1,
		)
		ok = bool(
			not list_left
			and w26
			and w26.dt == "Stock Reconciliation"
			and int(w26.enabled or 0) == 1
			and w26.view == "Form"
			and upgrade.SR_FORM_SCRIPT_MARKER in (w26.script or "")
			and "has_role" in (w26.script or "")
			and "Import Stock Count" in (w26.script or "")
			and form_script
			and int(form_script.enabled or 0) == 1
			and form_script.view == "Form"
		)
		check(
			"client_scripts",
			ok,
			f"w25_left={list_left}, "
			f"w26={ {k: w26.get(k) for k in ('dt', 'enabled', 'view')} if w26 else None }, "
			f"form23={dict(form_script) if form_script else None}",
		)
	except Exception as e:
		check("client_scripts", False, f"{type(e).__name__}: {str(e)[:200]}")

	# --- Fixture: item1 (stok 24), item2 (stok 0), item3 (batch) ---
	try:
		item1 = _make_item(stock, alt, item_group, "Gate Item")
		item2 = _make_item(stock, alt, item_group, "Zero Item")
		item3 = _make_item(stock, alt, item_group, "Batch Item")
		frappe.db.set_value("Item", item3.name, "has_batch_no", 1)
		_make_bin(item2.name, warehouse)
		_make_bin(item3.name, warehouse)
		frappe.db.commit()
		check(
			"fixture_item",
			True,
			f"item1={item1.name}, item2={item2.name}, item3(batch)={item3.name}",
		)
	except Exception as e:
		check("fixture_item", False, f"{type(e).__name__}: {str(e)[:200]}")
		raise GateAborted()

	# --- 2. Modul + endpoint (smoke decode base64/data-URL via endpoint, murni) ---
	try:
		def methods_for(module, name):
			# dict allowed_http_methods ber-key OBJECT fungsi — bandingkan via
			# module+name agar tahan bila whitelist membungkus fungsi
			return next(
				(
					m
					for fn, m in frappe.allowed_http_methods_for_whitelisted_func.items()
					if getattr(fn, "__module__", "") == module
					and getattr(fn, "__name__", "") == name
				),
				None,
			)

		header = ",".join(sr_import.TEMPLATE_HEADERS)
		tiny = f"{header}\n{item1.name},X,{warehouse},{alt},2,3,1200".encode()
		res = sr_import.upload_stock_count(
			filename="x.csv",
			data="data:text/csv;base64," + base64.b64encode(tiny).decode(),
			company=company,
		)
		g = (res.get("rows") or [{}])[0]
		ok = bool(
			methods_for(sr_import.__name__, "upload_stock_count") == ["POST"]
			and methods_for(sr_import.__name__, "download_stock_count_template") == ["POST"]
			and "has_permission" in inspect.getsource(sr_import.upload_stock_count)
			and "has_permission" in inspect.getsource(sr_import.download_stock_count_template)
			and "base64" in inspect.getsource(sr_import.upload_stock_count)
			and res.get("company") == company
			and res.get("added") == 1
			and abs(flt(g.get("qty")) - 36) < 1e-4  # 3 × 12
		)
		check(
			"module_endpoints",
			ok,
			f"sr_import: download/upload whitelisted POST + has_permission create SR + "
			f"data-URL decode; smoke added={res.get('added')}, qty={g.get('qty')}",
		)
	except Exception as e:
		frappe.db.rollback()
		check("module_endpoints", False, f"{type(e).__name__}: {str(e)[:200]}")

	if not _seed_stock(check, item1.name, warehouse, company):
		raise GateAborted()

	# --- 3. build_count_rows: prefill + toggle + eksklusi batch ---
	try:
		rows = sr_import.build_count_rows(warehouse, False)
		r1 = next((r for r in rows if r["item_code"] == item1.name), None)
		ok = bool(
			r1
			and r1["uom"] == alt
			and abs(flt(r1["current_qty"]) - 2) < 1e-6
			and not any(r["item_code"] == item2.name for r in rows)  # stok 0 tanpa toggle
			and not any(r["item_code"] == item3.name for r in rows)  # batch dikecualikan
		)
		rows_all = sr_import.build_count_rows(warehouse, True)
		r2 = next((r for r in rows_all if r["item_code"] == item2.name), None)
		ok = ok and bool(r2 and abs(flt(r2["current_qty"])) < 1e-6)
		ok = ok and not any(r["item_code"] == item3.name for r in rows_all)
		check(
			"template_rows",
			ok,
			f"r1={r1}, r2(toggle)={r2}, batch3 excluded={'yes' if not any(r['item_code'] == item3.name for r in rows_all) else 'NO'}",
		)
	except Exception as e:
		check("template_rows", False, f"{type(e).__name__}: {str(e)[:200]}")

	# --- 4. Roundtrip file (xlsx + csv + varian locale) ---
	_check_roundtrip(check, item1, item2, alt, warehouse)

	# --- 5. import_rows: validate → baris grid → SR dibangun klien-style ---
	draft_name = None
	posting_date = nowdate()
	posting_time = nowtime()
	try:
		from openpyxl import load_workbook

		rows = sr_import.build_count_rows(warehouse, True)
		xlsx_bytes = sr_import.make_template_bytes(rows, "xlsx")
		wb = load_workbook(io.BytesIO(xlsx_bytes))
		ws = wb.active
		idx = next(i for i, r in enumerate(rows) if r["item_code"] == item1.name)
		ws.cell(row=idx + 2, column=6, value=3)  # Counted Qty (as per UOM)
		ws.cell(row=idx + 2, column=7, value=1200)  # Valuation Rate (as per UOM)
		buf = io.BytesIO()
		wb.save(buf)
		parsed = sr_import.parse_count_content(buf.getvalue(), "count.xlsx")
		res = sr_import.validate_count_rows(parsed, company, posting_date, posting_time)
		g = (res.get("rows") or [{}])[0]
		grid_ok = bool(
			res.get("company") == company
			and res.get("added") == 1
			and res.get("skipped") >= 1  # baris item lain tanpa Counted Qty terskip
			and g.get("item_code") == item1.name
			and g.get("warehouse") == warehouse
			and g.get("stock_uom") == stock
			and g.get(SR_UOM_FIELD) == alt
			and abs(flt(g.get(SR_FACTOR_FIELD)) - 12) < 1e-6
			and abs(flt(g.get("current_qty")) - 24) < 1e-4
			and abs(flt(g.get(SR_BEFORE_FIELD)) - 2) < 1e-6
			and abs(flt(g.get(SR_AFTER_FIELD)) - 3) < 1e-6
			and abs(flt(g.get(SR_DIFF_FIELD)) - 1) < 1e-6
			and abs(flt(g.get("qty")) - 36) < 1e-4  # 3 × 12
			and abs(flt(g.get("valuation_rate")) - 100) < 1e-4  # 1200 ÷ 12
		)
		# Bangun SR persis seperti klien (fill grid dengan hasil endpoint, lalu
		# Save) — hook W23 yang mengonversi native qty/rate saat simpan.
		doc = frappe.get_doc(
			{
				"doctype": "Stock Reconciliation",
				"purpose": "Stock Reconciliation",
				"company": company,
				"posting_date": posting_date,
				"posting_time": posting_time,
				"items": res.get("rows") or [],
			}
		).insert()
		row = frappe.db.get_value(
			"Stock Reconciliation Item",
			{"parent": doc.name},
			[
				"qty",
				"valuation_rate",
				SR_FACTOR_FIELD,
				SR_AFTER_FIELD,
				SR_BEFORE_FIELD,
				SR_DIFF_FIELD,
				SR_RATE_FIELD,
			],
			as_dict=1,
		)
		docstatus = frappe.db.get_value("Stock Reconciliation", doc.name, "docstatus")
		ok = grid_ok and bool(
			row
			and abs(flt(row.qty) - 36) < 1e-4
			and abs(flt(row.valuation_rate) - 100) < 1e-4
			and abs(flt(row.get(SR_FACTOR_FIELD)) - 12) < 1e-6
			and abs(flt(row.get(SR_AFTER_FIELD)) - 3) < 1e-6
			and abs(flt(row.get(SR_BEFORE_FIELD)) - 2) < 1e-6
			and abs(flt(row.get(SR_DIFF_FIELD)) - 1) < 1e-6
			and abs(flt(row.get(SR_RATE_FIELD)) - 1200) < 1e-6
			and docstatus == 0
		)
		TRACKED["sr"].append(doc.name)
		draft_name = doc.name
		frappe.db.commit()
		check(
			"import_rows",
			ok,
			f"grid added={res.get('added')}, skipped={res.get('skipped')}, g={g}; "
			f"sr={doc.name}, row={dict(row) if row else None}, docstatus={docstatus}",
		)
	except Exception as e:
		frappe.db.rollback()
		check("import_rows", False, f"{type(e).__name__}: {str(e)[:200]}")

	# --- 6. Semua baris cocok ledger → tanpa dokumen ---
	try:
		before_srs = set(
			frappe.get_all(
				"Stock Reconciliation Item", filters={"item_code": item1.name}, pluck="parent"
			)
		)
		res = sr_import.validate_count_rows(
			{
				"rows": [
					{
						"row_no": 2,
						"item_code": item1.name,
						"warehouse": warehouse,
						"uom": alt,
						"counted": 2.0,
						"rate": None,
					}
				],
				"skipped": 0,
				"warnings": [],
			},
			company,
		)
		after_srs = set(
			frappe.get_all(
				"Stock Reconciliation Item", filters={"item_code": item1.name}, pluck="parent"
			)
		)
		ok = bool(
			res.get("added") == 0
			and res.get("rows") == []
			and res.get("matched") == 1
			and "match" in (res.get("message") or "").lower()
			and before_srs == after_srs
		)
		check(
			"matched_all_no_rows",
			ok,
			f"res={{'added': {res.get('added')}, 'matched': {res.get('matched')}, "
			f"'message': {res.get('message')!r}}}, srs before==after: {before_srs == after_srs}",
		)
	except Exception as e:
		check("matched_all_no_rows", False, f"{type(e).__name__}: {str(e)[:200]}")

	# --- 7. Error parse + error validate atomik ---
	_check_parse_errors(check, item1, alt, warehouse)
	_check_atomic_errors(check, item1, alt, warehouse, company)

	# --- 8. Permission user tanpa role (upload_stock_count langsung) ---
	try:
		email = (PREFIX + "-" + frappe.generate_hash(length=6) + "@example.com").lower()
		frappe.get_doc(
			{
				"doctype": "User",
				"email": email,
				"first_name": PREFIX,
				"last_name": "Gate",
				"user_type": "System User",
				"send_welcome_email": 0,
			}
		).insert()
		TRACKED["user"].append(email)
		frappe.db.commit()
		denied = False
		frappe.set_user(email)
		try:
			sr_import.upload_stock_count(filename="x.csv", data="aGk=", company=company)
		except frappe.PermissionError:
			denied = True
		finally:
			frappe.set_user("Administrator")
		check(
			"permission_gate",
			denied,
			f"user={email} tanpa role → upload_stock_count ditolak: {denied}",
		)
	except Exception as e:
		frappe.set_user("Administrator")
		check("permission_gate", False, f"{type(e).__name__}: {str(e)[:200]}")

	# --- 9. Item stok-0 tanpa rate → warning + baris tetap dikembalikan ---
	try:
		res = sr_import.validate_count_rows(
			{
				"rows": [
					{
						"row_no": 2,
						"item_code": item2.name,
						"warehouse": warehouse,
						"uom": alt,
						"counted": 5.0,
						"rate": None,
					}
				],
				"skipped": 0,
				"warnings": [],
			},
			company,
		)
		g = (res.get("rows") or [{}])[0]
		ok = bool(
			res.get("added") == 1
			and len(res.get("rows") or []) == 1
			and any("Valuation Rate" in w for w in res.get("warnings") or [])
			and abs(flt(g.get("qty")) - 60) < 1e-4  # 5 × 12
		)
		check(
			"zero_ledger_warning",
			ok,
			f"res={{'added': {res.get('added')}, 'warnings': {res.get('warnings')}}}, qty={g.get('qty')}",
		)
	except Exception as e:
		check("zero_ledger_warning", False, f"{type(e).__name__}: {str(e)[:200]}")

	# --- 10. Submit draft hasil import → SLE ---
	try:
		if not draft_name:
			check("submit_draft_sle", False, "draft dari import_rows tidak tersedia")
		else:
			doc = frappe.get_doc("Stock Reconciliation", draft_name)
			doc.submit()
			frappe.db.commit()
			sle = frappe.db.get_value(
				"Stock Ledger Entry",
				{"voucher_type": "Stock Reconciliation", "voucher_no": draft_name, "is_cancelled": 0},
				["qty_after_transaction", "valuation_rate"],
				as_dict=1,
			)
			ok = bool(
				sle
				and abs(flt(sle.qty_after_transaction) - 36) < 1e-4
				and abs(flt(sle.valuation_rate) - 100) < 1e-4
			)
			check(
				"submit_draft_sle",
				ok,
				f"sr={draft_name}, sle={dict(sle) if sle else None}, want qty=36 @100",
			)
	except Exception as e:
		frappe.db.rollback()
		check("submit_draft_sle", False, f"{type(e).__name__}: {str(e)[:200]}")

	# --- 11. apply() idempoten: pensiunan W25 no-op, W26 tetap hidup ---
	try:
		upgrade.apply()
		upgrade.apply()
		list_left = frappe.get_all(
			"Client Script", filters={"script": ("like", "%" + MARKER + "%")}, pluck="name"
		)
		w26 = frappe.db.get_value(
			"Client Script",
			"warehouse-app-w26-stock-count-import",
			["enabled", "view"],
			as_dict=1,
		)
		ok = bool(not list_left and w26 and int(w26.enabled or 0) == 1 and w26.view == "Form")
		check(
			"upgrade_idempotent",
			ok,
			f"apply()×2 ok, w25_left={list_left}, w26={dict(w26) if w26 else None}",
		)
	except Exception as e:
		check("upgrade_idempotent", False, f"{type(e).__name__}: {str(e)[:200]}")


def _check_roundtrip(check, item1, item2, alt, warehouse):
	"""4+6: xlsx & csv roundtrip (edit sel via openpyxl / csv module), varian
	locale delimiter ';' + desimal koma + grouping titik, dan baris tanpa
	count terskip parser."""
	try:
		from openpyxl import load_workbook

		rows = sr_import.build_count_rows(warehouse, True)
		xlsx_bytes = sr_import.make_template_bytes(rows, "xlsx")
		wb = load_workbook(io.BytesIO(xlsx_bytes))
		ws = wb.active
		header = [c.value for c in next(ws.iter_rows(max_row=1))]
		idx = next(i for i, r in enumerate(rows) if r["item_code"] == item1.name)
		sheet_row = idx + 2
		ws.cell(row=sheet_row, column=6, value=3)  # Counted Qty (as per UOM)
		ws.cell(row=sheet_row, column=7, value=1200)  # Valuation Rate (as per UOM)
		buf = io.BytesIO()
		wb.save(buf)

		parsed_x = sr_import.parse_count_content(buf.getvalue(), "count.xlsx")
		px = next((r for r in parsed_x["rows"] if r["item_code"] == item1.name), None)
		ok_x = bool(
			header[:1] == [sr_import.H_ITEM]
			and px
			and abs(flt(px["counted"]) - 3) < 1e-9
			and abs(flt(px["rate"]) - 1200) < 1e-9
			and px["uom"] == alt
		)

		csv_bytes = sr_import.make_template_bytes(rows, "csv")
		lines = list(csv.reader(io.StringIO(csv_bytes.decode("utf-8-sig"))))
		cidx = next(i for i, r in enumerate(lines) if r and r[0] == item1.name)
		lines[cidx][5] = "3"
		lines[cidx][6] = "1200"
		# baris item2 tanpa count → harus terskip parser (stat skipped)
		out = io.StringIO()
		csv.writer(out).writerows(lines)
		parsed_c = sr_import.parse_count_content(out.getvalue().encode("utf-8"), "count.csv")
		pc = next((r for r in parsed_c["rows"] if r["item_code"] == item1.name), None)
		ok_c = bool(
			pc
			and abs(flt(pc["counted"]) - 3) < 1e-9
			and abs(flt(pc["rate"]) - 1200) < 1e-9
			and not any(r["item_code"] == item2.name for r in parsed_c["rows"])
			and parsed_c["skipped"] >= 1
		)

		# varian locale Excel ID: delimiter ';' + "3,5" + grouping "1.500"
		semi = "\n".join(
			[
				";".join(sr_import.TEMPLATE_HEADERS),
				f"{item1.name};X;{warehouse};{alt};2;3,5;1200,5",
				f"{item2.name};X;{warehouse};{alt};0;1.500;",
			]
		)
		parsed_s = sr_import.parse_count_content(semi.encode("utf-8"), "count.csv")
		ps1 = next((r for r in parsed_s["rows"] if r["item_code"] == item1.name), None)
		ps2 = next((r for r in parsed_s["rows"] if r["item_code"] == item2.name), None)
		ok_s = bool(
			ps1
			and abs(flt(ps1["counted"]) - 3.5) < 1e-9
			and abs(flt(ps1["rate"]) - 1200.5) < 1e-9
			and ps2
			and abs(flt(ps2["counted"]) - 1500) < 1e-9
		)

		# dialek koma: titik selalu desimal — "1.500" = 1.5, BUKAN grouping 1500
		comma_header = ",".join(sr_import.TEMPLATE_HEADERS)
		parsed_g = sr_import.parse_count_content(
			f"{comma_header}\n{item2.name},X,{warehouse},{alt},0,1.500,".encode(), "count.csv"
		)
		pg = next((r for r in parsed_g["rows"] if r["item_code"] == item2.name), None)
		ok_g = bool(pg and abs(flt(pg["counted"]) - 1.5) < 1e-9)

		check(
			"roundtrip_files",
			ok_x and ok_c and ok_s and ok_g,
			f"xlsx={ok_x} csv={ok_c} locale_semi={ok_s} comma_decimal={ok_g}; "
			f"px={px}, pc={pc}, ps1={ps1}, ps2={ps2}, pg={pg}, csv_skipped={parsed_c.get('skipped')}",
		)
	except Exception as e:
		check("roundtrip_files", False, f"{type(e).__name__}: {str(e)[:250]}")


def _parse_err(content, filename):
	try:
		sr_import.parse_count_content(content, filename)
		return None
	except frappe.ValidationError as e:
		frappe.db.rollback()
		return str(e)
	except frappe.PermissionError:
		raise


def _check_parse_errors(check, item1, alt, warehouse):
	try:
		header = ",".join(sr_import.TEMPLATE_HEADERS)
		results = {
			"bad_ext": _parse_err(b"a,b\n1,2", "count.txt"),
			"missing_col": _parse_err(b"A,B\n1,2", "count.csv"),
			"bad_xlsx": _parse_err(b"bukan xlsx sama sekali", "count.xlsx"),
			"rate_zero": _parse_err(
				f"{header}\n{item1.name},X,{warehouse},{alt},2,3,0".encode(), "count.csv"
			),
			"garbage_number": _parse_err(
				f"{header}\n{item1.name},X,{warehouse},{alt},2,3,abc".encode(), "count.csv"
			),
			"mixed_sep": _parse_err(
				f"{header}\n{item1.name},X,{warehouse},{alt},2,\"1,5.5\",1200".encode(), "count.csv"
			),
		}
		ok = bool(
			results["bad_ext"]
			and "Only .xlsx" in results["bad_ext"]
			and results["missing_col"]
			and "Missing required column" in results["missing_col"]
			and results["bad_xlsx"]
			and "not a valid .xlsx" in results["bad_xlsx"]
			and results["rate_zero"]
			and "Row 2" in results["rate_zero"]
			and "leave the cell blank" in results["rate_zero"]
			and results["garbage_number"]
			and "not a recognized number" in results["garbage_number"]
			and results["mixed_sep"]
			and "mixes" in results["mixed_sep"]
		)
		check(
			"parse_errors",
			ok,
			"; ".join(f"{k}={'OK' if v else 'NO-THROW'}" for k, v in results.items())
			+ f" | rate_zero[:120]={results['rate_zero'][:120] if results['rate_zero'] else None}",
		)
	except Exception as e:
		check("parse_errors", False, f"{type(e).__name__}: {str(e)[:200]}")


def _check_atomic_errors(check, item1, alt, warehouse, company):
	try:
		def rowdict(row_no, item, counted, rate=None, uom=None):
			return {
				"row_no": row_no,
				"item_code": item,
				"warehouse": warehouse,
				"uom": uom if uom is not None else alt,
				"counted": counted,
				"rate": rate,
			}

		cases = [
			("unknown_item", [rowdict(2, PREFIX + "-NOPE", 1)], "not found", 2, company),
			# duplikat dilaporkan pada kemunculan KEDUA (Row 3)
			("duplicate", [rowdict(2, item1.name, 1), rowdict(3, item1.name, 2)], "more than once", 3, company),
			("uom_empty", [rowdict(2, item1.name, 1, uom="")], "UOM is required", 2, company),
			("negative", [rowdict(2, item1.name, -5)], "negative", 2, company),
			("rate_zero", [rowdict(2, item1.name, 1, rate=0)], "leave the cell blank", 2, company),
			# Mismatch TANPA membuat Company (insert Company = berat, membangun
			# chart of accounts): param company sengaja TIDAK ADA di site,
			# gudangnya milik company JURI → cabang "belongs to company" terpicu
			# dan pesan memuat nama company milik gudang.
			(
				"company_mismatch",
				[rowdict(2, item1.name, 1)],
				f"belongs to company|{company}",
				2,
				PREFIX + "-NOCO",
			),
		]
		results = {}
		for name, rows, needle, want_row, use_company in cases:
			try:
				sr_import.validate_count_rows({"rows": rows, "skipped": 0, "warnings": []}, use_company)
				results[name] = "NO-THROW"
			except frappe.ValidationError as e:
				msg = str(e)
				results[name] = (
					"OK"
					if (f"Row {want_row}" in msg and all(x in msg for x in needle.split("|")))
					else f"WRONG-MSG: {msg[:120]}"
				)
				frappe.db.rollback()
		persisted = frappe.db.count("Stock Reconciliation Item", {"item_code": item1.name})
		ok = all(v == "OK" for v in results.values())
		check(
			"atomic_errors",
			ok,
			f"{results}; baris SR item1 tersisa={persisted} (hanya draft import_rows)",
		)
	except Exception as e:
		frappe.db.rollback()
		check("atomic_errors", False, f"{type(e).__name__}: {str(e)[:200]}")


def _teardown(check):
	try:
		frappe.set_user("Administrator")
	except Exception:
		pass
	sweep_ok, sweep_evidence = _sweep()
	residue = count_residue(
		PREFIX, {"Stock Reconciliation": TRACKED["sr"], "Stock Entry": TRACKED["se"]}
	)
	zero = all(v == 0 for v in residue.values())
	users_left = [u for u in TRACKED["user"] if frappe.db.exists("User", u)]
	check(
		"teardown",
		sweep_ok and zero and not users_left,
		f"{sweep_evidence}; residu {PREFIX}={residue} (SR tracked={len(TRACKED['sr'])}, "
		f"SE tracked={len(TRACKED['se'])}), users left={users_left}",
	)
	TRACKED.update({"sr": [], "se": [], "item": [], "user": []})


def _sweep():
	"""Bersihkan residu W25/W26: SR/SE ter-track + via item prefix, user fixture
	(+ Sessions/Activity Log), SLE/Bin/Repost Item Valuation per item (delete
	tabel repost di-guard safe() — tabelnya tak selalu ada di site), Item,
	Version. Idempoten."""
	errors = []

	def safe(label, fn):
		try:
			fn()
		except Exception as e:
			errors.append(f"{label}: {type(e).__name__}: {str(e)[:120]}")

	def cancel_delete(doctype, name):
		if not frappe.db.exists(doctype, name):
			return
		doc = frappe.get_doc(doctype, name)
		if doc.docstatus == 1:
			doc.flags.ignore_permissions = True
			doc.cancel()
		frappe.delete_doc(doctype, name, force=1, ignore_missing=True)

	items = frappe.get_all(
		"Item",
		or_filters=[
			["Item", "name", "like", PREFIX + "%"],
			["Item", "item_code", "like", PREFIX + "%"],
			["Item", "item_name", "like", PREFIX + "%"],
		],
		pluck="name",
	)
	doc_names = set(TRACKED["sr"]) | set(TRACKED["se"])
	if items:
		for dt in ("Stock Reconciliation Item", "Stock Entry Detail"):
			doc_names.update(
				frappe.get_all(dt, filters={"item_code": ("in", items)}, pluck="parent")
			)
	for name in doc_names:
		if frappe.db.exists("Stock Reconciliation", name):
			safe(f"SR {name}", lambda n=name: cancel_delete("Stock Reconciliation", n))
			safe(
				f"GL SR {name}",
				lambda n=name: frappe.db.delete(
					"GL Entry", {"voucher_type": "Stock Reconciliation", "voucher_no": n}
				),
			)
		elif frappe.db.exists("Stock Entry", name):
			safe(f"SE {name}", lambda n=name: cancel_delete("Stock Entry", n))
			safe(
				f"GL SE {name}",
				lambda n=name: frappe.db.delete(
					"GL Entry", {"voucher_type": "Stock Entry", "voucher_no": n}
				),
			)
	# SLE/Bin per item (jaga-jaga bila cancel tak sempat membersihkan)
	safe(
		"SLE",
		lambda: frappe.db.delete("Stock Ledger Entry", {"item_code": ("in", items or [""])})
		if items
		else None,
	)
	safe(
		"Repost",
		lambda: frappe.db.delete("Repost Item Valuation", {"item_code": ("in", items or [""])})
		if items
		else None,
	)
	safe(
		"Bin", lambda: frappe.db.delete("Bin", {"item_code": ("in", items or [""])}) if items else None
	)
	# Komit per fase: error MariaDB kelas "record has changed" (1020) membatalkan
	# SELURUH transaksi berjalan — tanpa komit antar fase, kegagalan di fase user
	# mengembalikan pembersihan dokumen yang sudah dilakukan di atasnya.
	frappe.db.commit()
	# User fixture + jejak sesi/aktivitas. Ditemukan via TRACKED maupun pola
	# prefix — residu run sebelumnya yang crash tidak ada di TRACKED run ini.
	# Contact dihapus dulu di txn terpisah: delete_doc User menyentuh Contact
	# terkait dan memicu QueryDeadlockError tabContact bila Contact itu basi
	# di transaksi yang sama.
	users = set(TRACKED["user"])
	users.update(
		frappe.get_all(
			"User",
			or_filters=[
				["User", "name", "like", PREFIX + "%"],
				["User", "first_name", "like", PREFIX + "%"],
			],
			pluck="name",
		)
	)
	users.discard("Administrator")
	for email in users:
		safe(f"Contact {email}", lambda e=email: _drop_contact(e))
	safe("Contact orphan", _drop_orphan_contacts)
	frappe.db.commit()
	for email in users:
		safe(f"User {email}", lambda e=email: frappe.delete_doc("User", e, force=1))
		safe(f"Sessions {email}", lambda e=email: frappe.db.delete("Sessions", {"user": e}))
		safe(f"Activity {email}", lambda e=email: frappe.db.delete("Activity Log", {"user": e}))
	frappe.db.commit()
	# Item
	for name in items:
		safe(f"Item {name}", lambda n=name: frappe.delete_doc("Item", n, force=1, ignore_missing=True))
	safe("Version", lambda: frappe.db.delete("Version", {"docname": ("like", PREFIX + "%")}))
	frappe.db.commit()

	return (not errors), ("bersih" if not errors else "; ".join(errors[:5]))


def _drop_contact(email):
	for name in frappe.get_all("Contact", filters={"email_id": email}, pluck="name"):
		frappe.delete_doc("Contact", name, force=1, ignore_missing=True)


def _drop_orphan_contacts():
	"""Contact sisa run crash (user-nya sudah tidak ada, tak bisa dicari via email)."""
	for name in frappe.get_all(
		"Contact",
		or_filters=[
			["Contact", "first_name", "like", PREFIX + "%"],
			["Contact", "last_name", "like", PREFIX + "%"],
		],
		pluck="name",
	):
		frappe.delete_doc("Contact", name, force=1, ignore_missing=True)


def _pick_uoms():
	stock = next((u for u in ("Nos", "Unit", "Kg") if frappe.db.exists("UOM", u)), None)
	alt = next(
		(u for u in ("Pack", "Box", "Bag") if u != stock and frappe.db.exists("UOM", u)), None
	)
	invalid = next(
		(u for u in ("Kg", "Gram", "Meter") if u not in (stock, alt) and frappe.db.exists("UOM", u)),
		None,
	)
	return stock, alt, invalid


def _pick_item_group():
	for group in ("Bread", "Products"):
		if frappe.db.exists("Item Group", group) and not frappe.db.get_value("Item Group", group, "is_group"):
			return group
	return frappe.db.get_value("Item Group", {"is_group": 0}, "name", order_by="name")
